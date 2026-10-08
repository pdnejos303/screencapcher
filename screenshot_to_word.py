import datetime
import os
from docx import Document
from docx.shared import Inches
import keyboard
import pyautogui

# ตั้งชื่อไฟล์ Word
doc_name = "my_report.docx"

# เช็คว่ามีไฟล์เดิมอยู่หรือไม่ เพื่อป้องกันการเขียนทับของเก่าทั้งหมด
if os.path.exists(doc_name):
    try:
        doc = Document(doc_name)
        print(f"โหลดไฟล์ {doc_name} เดิมมาใช้งานแล้ว")
    except Exception as e:
        print(f"ไม่สามารถเปิดไฟล์ {doc_name} ได้ กรุณาปิดโปรแกรม Word ก่อนเริ่มทำงาน: {e}")
        exit()
else:
    doc = Document()
    doc.add_heading("Automation Screenshot Report", 0)
    print("สร้างไฟล์ Word ใหม่เรียบร้อย")

def capture_and_save():
    try:
        # ตั้งชื่อไฟล์รูปจากเวลาปัจจุบัน
        filename = f"screenshot_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.png"

        # แคปหน้าจอแล้วบันทึกชั่วคราว
        screenshot = pyautogui.screenshot()
        screenshot.save(filename)

        # เอาภาพใส่ Word
        doc.add_paragraph(f"Captured at: {datetime.datetime.now()}")
        doc.add_picture(filename, width=Inches(5.5))  # กำหนดขนาดรูป
        
        # บันทึกไฟล์ Word
        doc.save(doc_name)
        print(f"\n[สำเร็จ] ถ่ายภาพและบันทึกลง {doc_name} เรียบร้อยแล้ว!")
        
        # ลบรูปภาพชั่วคราวทิ้ง (เพื่อไม่ให้รกเครื่อง)
        if os.path.exists(filename):
            os.remove(filename)
            
        # เปิดไฟล์ Word ขึ้นมาให้เห็นทันที
        os.startfile(doc_name)
        print("เปิดไฟล์ Word ให้ดูแล้ว (หมายเหตุ: กรุณาปิดหน้าต่าง Word ก่อนกด Shift+S เพื่อถ่ายภาพครั้งถัดไป)")
        
    except PermissionError:
        print("\n[แจ้งเตือน] ไม่สามารถบันทึกไฟล์ได้เนื่องจากไฟล์ Word ยังเปิดค้างอยู่! กรุณาปิดหน้าต่าง Word ก่อนกด Shift+S อีกครั้ง")
    except Exception as e:
        print(f"\n[เกิดข้อผิดพลาด] {e}")

# ตั้งค่าปุ่มลัด กด Shift+S ให้แคป
keyboard.add_hotkey("shift+s", capture_and_save)

print("="*50)
print("โปรแกรมเริ่มทำงานแล้ว!")
print(" - กดปุ่ม 'Shift + S' เพื่อถ่ายภาพหน้าจอและบันทึกลง Word")
print(" - กดปุ่ม 'ESC' เพื่อปิดโปรแกรมนี้")
print("="*50)

# รอการกดปุ่ม (โปรแกรมจะทำงานจนกว่าจะกด esc)
keyboard.wait("esc")
print("ปิดโปรแกรมเรียบร้อยแล้ว")
