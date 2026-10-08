import datetime
import os
import keyboard
import sys
from mss import mss
from mss.tools import to_png
from docx import Document
from docx.shared import Inches

sys.stdout.reconfigure(line_buffering=True)

print("กำลังเตรียมระบบ...")
sct = mss()
monitors = sct.monitors

# เลือกจอที่ 2 อัตโนมัติถ้ามีหลายจอ
if len(monitors) > 2:
    selected_monitor_idx = 2
    print("เลือกแคปภาพจาก 'จอที่ 2' อัตโนมัติ")
else:
    selected_monitor_idx = 1
    print("เลือกแคปภาพจาก 'จอหลัก'")

monitor_to_capture = monitors[selected_monitor_idx]
current_dir = os.path.abspath(os.path.dirname(__file__))
doc_name = "MyReport.docx"
doc_path = os.path.join(current_dir, doc_name)

# สร้างไฟล์รอไว้ถ้ายังไม่มี
if not os.path.exists(doc_path):
    doc = Document()
    doc.add_heading("รายงานภาพหน้าจออัตโนมัติ", level=1)
    doc.save(doc_path)

def capture_and_save():
    print(f"\n📸 กำลังแคปหน้าจอ...")
    
    # 1. แคปและบันทึกภาพชั่วคราว
    image_filename = os.path.join(current_dir, "screenshot_temp.png")
    sct_img = sct.grab(monitor_to_capture)
    to_png(sct_img.rgb, sct_img.size, output=image_filename)
    
    # 2. นำภาพมาใส่ Word
    try:
        doc = Document(doc_path)
        doc.add_paragraph(f"บันทึกภาพเมื่อเวลา: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_picture(image_filename, width=Inches(6.0))
        
        # 3. บันทึกทับไฟล์เดิม
        doc.save(doc_path)
        print(f"✅ บันทึกภาพลงไฟล์ Word สำเร็จ!")
        
        # 4. สั่งเปิดไฟล์ Word ขึ้นมาโชว์ทันที
        os.startfile(doc_path)
        print("📂 เปิดไฟล์ Word ให้ดูแล้วครับ! (สำคัญ: ต้องปิดหน้าต่าง Word นี้ก่อน ถึงจะกด Shift+S ถ่ายรูปครั้งต่อไปได้)")
        
    except PermissionError:
        print(f"❌ แจ้งเตือน: เซฟไม่ได้เพราะไฟล์ Word เปิดค้างอยู่! กรุณาปิดโปรแกรม Word ก่อนกด Shift+S อีกครั้ง")
    except Exception as e:
        print(f"❌ เกิดข้อผิดพลาด: {e}")
        
    # ลบไฟล์ภาพทิ้ง
    try:
        os.remove(image_filename)
    except:
        pass

keyboard.add_hotkey("shift+s", capture_and_save)

print("="*70)
print(f"✅ โปรแกรมพร้อมทำงานแล้ว! (ใช้ python-docx ตามตัวอย่างที่คุณส่งมา)")
print(f" - กดปุ่ม 'Shift + S' ➔ โปรแกรมจะแคปภาพ ➔ เซฟลง Word ➔ แล้ว Word จะเด้งเปิดขึ้นมาโชว์ทันที!")
print(f" - กดปุ่ม 'ESC' เพื่อปิดโปรแกรม")
print("="*70)

keyboard.wait("esc")
